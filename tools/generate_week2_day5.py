from pathlib import Path
import base64, shutil, zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'week2' / 'day5' / 'downloads' / 'Day05_CALCULATE_VISUALIZE.zip'
WORK = ROOT / 'week2' / 'day5' / '_generated_day5'
PBIX_PARTS = ROOT / 'tools' / 'day5_pbix_parts'

TEAL='099999'
NAVY='13233B'
GOLD='F2C811'
WHITE='FFFFFF'
LIGHT='F3F7F7'
MID='D8E7E7'
TEXT='1D2A33'
MUTED='61717C'
RED='C84630'
GREEN='1F8A70'

thin = Side(style='thin', color='D8E2E5')

def style_sheet(ws, widths=None):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A5'
    if widths:
        for i,w in enumerate(widths,1):
            ws.column_dimensions[get_column_letter(i)].width=w

def title_block(ws, title, subtitle):
    ws.merge_cells('A1:F1')
    c=ws['A1']; c.value=title; c.font=Font(name='Almarai',size=20,bold=True,color=WHITE)
    c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='left',vertical='center')
    ws.row_dimensions[1].height=34
    ws.merge_cells('A2:F2')
    c=ws['A2']; c.value=subtitle; c.font=Font(name='Almarai',size=11,color=MUTED)
    c.alignment=Alignment(horizontal='left',vertical='center')
    ws.merge_cells('A3:F3')
    c=ws['A3']; c.value='X Academy • Power BI Specialist • Day 5'; c.font=Font(name='Almarai',size=10,bold=True,color=TEAL)
    ws.row_dimensions[3].height=22

def add_table(ws, start_row, headers, rows, widths=None):
    for col,h in enumerate(headers,1):
        c=ws.cell(start_row,col,h)
        c.font=Font(name='Almarai',bold=True,color=WHITE)
        c.fill=PatternFill('solid',fgColor=TEAL)
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        c.border=Border(bottom=thin)
    for r_idx,row in enumerate(rows,start_row+1):
        fill = LIGHT if (r_idx-start_row)%2==1 else WHITE
        for col,val in enumerate(row,1):
            c=ws.cell(r_idx,col,val)
            c.font=Font(name='Almarai',size=10,color=TEXT)
            c.fill=PatternFill('solid',fgColor=fill)
            c.alignment=Alignment(vertical='top',wrap_text=True)
            c.border=Border(bottom=thin)
        ws.row_dimensions[r_idx].height=36
    if widths:
        for i,w in enumerate(widths,1):
            ws.column_dimensions[get_column_letter(i)].width=w

def new_wb(title, subtitle):
    wb=Workbook()
    ws=wb.active
    ws.title='Guide'
    title_block(ws,title,subtitle)
    style_sheet(ws)
    return wb,ws

def save_power_query(path):
    wb,ws=new_wb('POWER QUERY CALCULATIONS','نبني 3 أعمدة جديدة داخل FactSales قبل الانتقال إلى DAX.')
    rows=[
      ['01','Gross Sales','قيمة البيع قبل الخصم','[Quantity] * [Unit Price]','Decimal Number','لا تحسب Total هنا؛ هذا Row-level calculation.'],
      ['02','Discount Amount','قيمة الخصم الفعلية','[Gross Sales] - [Revenue]','Decimal Number','Revenue هو المبلغ النهائي بعد الخصم في بيانات Nova.'],
      ['03','Discount Status','تصنيف الصف حسب وجود خصم','if [Discount %] > 0 then "Discounted" else "Full Price"','Text','سنستخدمه لاحقا داخل Slicer / Chart.']
    ]
    add_table(ws,5,['Step','New Column','Business Meaning','Custom Column Formula','Data Type','Trainer Note'],rows,[9,22,28,48,18,38])
    ws['A10']='الفكرة التي يجب أن تخرج بها'
    ws['A10'].font=Font(name='Almarai',size=12,bold=True,color=TEAL)
    ws.merge_cells('A11:F12')
    ws['A11']='Power Query يحسب هذه الأعمدة أثناء Refresh. أما KPI يتغير مع Filters داخل التقرير، فهذا سنبنيه كـ DAX Measure.'
    ws['A11'].alignment=Alignment(wrap_text=True,vertical='center')
    ws['A11'].font=Font(name='Almarai',size=12,bold=True,color=TEXT)
    ws['A11'].fill=PatternFill('solid',fgColor='E7F5F4')
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def save_dax(path):
    wb,ws=new_wb('FIRST DAX MEASURES','اكتب كل Measure بنفسك ثم اختبرها في Card قبل الانتقال للـVisual التالي.')
    rows=[
      ['01','Total Revenue','SUM','SUM(FactSales[Revenue])','Currency','KPI Card'],
      ['02','Total Quantity','SUM','SUM(FactSales[Quantity])','Whole Number','KPI Card'],
      ['03','Transactions','COUNTROWS','COUNTROWS(FactSales)','Whole Number','KPI Card'],
      ['04','Customers','DISTINCTCOUNT','DISTINCTCOUNT(FactSales[Customer ID])','Whole Number','KPI Card'],
      ['05','Average Transaction Value','DIVIDE','DIVIDE([Total Revenue], [Transactions])','Currency','Optional KPI'],
      ['06','Average Revenue per Customer','DIVIDE','DIVIDE([Total Revenue], [Customers])','Currency','Stretch Measure'],
      ['07','Total Discount Amount','SUM','SUM(FactSales[Discount Amount])','Currency','Stretch Measure']
    ]
    add_table(ws,5,['#','Measure','Function','Expression','Format','Use'],rows,[7,30,18,48,20,22])
    ws['A14']='Measure vs Power Query Column'
    ws['A14'].font=Font(name='Almarai',size=12,bold=True,color=TEAL)
    compare=[
      ['Power Query Column','قيمة لكل Row','Calculated during Refresh','Gross Sales / Discount Amount / Discount Status'],
      ['DAX Measure','نتيجة حسب Filter Context','Calculated when report is queried','Total Revenue / Customers / Transactions']
    ]
    add_table(ws,15,['Type','What it represents','When calculated','Examples'],compare,[25,32,32,55])
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def save_report(path):
    wb,ws=new_wb('SALES OVERVIEW — REPORT BLUEPRINT','أول Report حقيقي: KPIs + Trend + Category + Branch + Discount + Slicers.')
    layout=[
      ['KPI 1','Card','[Total Revenue]','','Format as Currency'],
      ['KPI 2','Card','[Customers]','','Whole Number'],
      ['KPI 3','Card','[Transactions]','','Whole Number'],
      ['KPI 4','Card','[Total Quantity]','','Whole Number'],
      ['Visual 1','Line Chart','DimDate[Month]','[Total Revenue]','Sort Month by DimDate[Month Number]'],
      ['Visual 2','Bar Chart','DimProduct[Category]','[Total Revenue]','Sort Descending'],
      ['Visual 3','Bar Chart','DimBranch[Branch Name]','[Total Revenue]','Visual Filter: Top 10'],
      ['Visual 4','Column / Donut','FactSales[Discount Status]','[Total Revenue]','Compare Discounted vs Full Price'],
      ['Slicer 1','Slicer','DimDate[Month]','',''],
      ['Slicer 2','Slicer','DimBranch[Branch Name]','',''],
      ['Slicer 3','Slicer','DimProduct[Category]','',''],
      ['Slicer 4','Slicer','FactSales[Discount Status]','','Optional']
    ]
    add_table(ws,5,['Element','Visual Type','Axis / Field','Values','Important Step'],layout,[14,22,34,25,42])
    ws['A19']='Formatting Checklist'
    ws['A19'].font=Font(name='Almarai',size=12,bold=True,color=TEAL)
    checklist=[
      ['01','Page title','NOVA SALES OVERVIEW'],
      ['02','Alignment','Cards same size and aligned'],
      ['03','Numbers','Currency + thousand separators + sensible decimals'],
      ['04','Titles','Rename visual titles to business language'],
      ['05','Noise','Remove unnecessary gridlines / labels'],
      ['06','Hierarchy','KPIs first → trends → breakdowns → slicers'],
      ['07','Test','Choose Category and confirm all KPIs / visuals respond']
    ]
    add_table(ws,20,['#','Check','Expected'],checklist,[7,24,70])
    # simple wireframe
    wire=wb.create_sheet('Wireframe')
    wire.sheet_view.showGridLines=False
    wire.merge_cells('B2:M3'); wire['B2']='NOVA SALES OVERVIEW'; wire['B2'].font=Font(name='Almarai',size=20,bold=True,color=WHITE); wire['B2'].fill=PatternFill('solid',fgColor=NAVY); wire['B2'].alignment=Alignment(horizontal='center',vertical='center')
    boxes=[('B5:D7','TOTAL REVENUE'),('E5:G7','CUSTOMERS'),('H5:J7','TRANSACTIONS'),('K5:M7','QUANTITY'),
           ('B9:G15','REVENUE TREND'),('H9:M15','REVENUE BY CATEGORY'),
           ('B17:G23','TOP BRANCHES'),('H17:M23','DISCOUNT ANALYSIS')]
    for rng,label in boxes:
        wire.merge_cells(rng); c=wire[rng.split(':')[0]]; c.value=label; c.font=Font(name='Almarai',bold=True,color=TEXT); c.fill=PatternFill('solid',fgColor='E7F5F4'); c.alignment=Alignment(horizontal='center',vertical='center'); c.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    wire.merge_cells('B25:M27'); wire['B25']='SLICERS: Month   |   Branch   |   Category   |   Discount Status'; wire['B25'].font=Font(name='Almarai',bold=True,color=TEAL); wire['B25'].alignment=Alignment(horizontal='center',vertical='center')
    for col in range(2,14): wire.column_dimensions[get_column_letter(col)].width=12
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def save_challenge(path):
    wb,ws=new_wb('DAY 5 — ANALYST CHALLENGE','ابنِ Sales Overview بدون تقليد خطوات المدرب واحدة واحدة.')
    req=[
      ['01','4 KPI Cards','Total Revenue • Customers • Transactions • Total Quantity'],
      ['02','1 Trend','Revenue by Month'],
      ['03','1 Category Analysis','Revenue by Category'],
      ['04','1 Branch Analysis','Top 10 Branches by Revenue'],
      ['05','1 Discount Analysis','Discount Status vs Revenue'],
      ['06','3+ Slicers','Month • Branch • Category (Discount Status optional)'],
      ['07','Formatting','Readable titles, aligned visuals, correct number formats']
    ]
    add_table(ws,5,['#','Requirement','Minimum Output'],req,[8,26,72])
    ws['A14']='Business Questions'
    ws['A14'].font=Font(name='Almarai',size=12,bold=True,color=TEAL)
    questions=[
      ['1','Which Category generated the highest Revenue?'],
      ['2','Which Branch appears strongest by Revenue?'],
      ['3','How much Revenue came from Discounted transactions?'],
      ['4','Choose one Category. What changed in Customers and Transactions?'],
      ['5','Why did the Measures change even though you did not edit the DAX?']
    ]
    add_table(ws,15,['#','Question'],questions,[8,100])
    ws.merge_cells('A22:F24')
    ws['A22']='Bridge to Day 6: نفس الـMeasure أعطى نتيجة مختلفة بعد اختيار Filter. غدا سنفهم لماذا — Filter Context + CALCULATE.'
    ws['A22'].font=Font(name='Almarai',size=13,bold=True,color=WHITE)
    ws['A22'].fill=PatternFill('solid',fgColor=TEAL)
    ws['A22'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def restore_start_pbix(path):
    parts = sorted(PBIX_PARTS.glob('part*.b64'))
    if not parts:
        raise FileNotFoundError(f'No PBIX base64 parts found in {PBIX_PARTS}')
    payload = ''.join(p.read_text(encoding='utf-8').strip() for p in parts)
    path.write_bytes(base64.b64decode(payload))

def build():
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)
    restore_start_pbix(WORK/'01_Nova_Day05_Start.pbix')

    (WORK/'00_START_HERE.txt').write_text(
      'NOVA DISTRIBUTION GROUP — DAY 05: CALCULATE & VISUALIZE\n\n'
      'IMPORTANT\n'
      'There is NO new dataset in Day 5. Start from the included Nova_Day05_Start.pbix, which is the Day 4 model supplied by the trainer.\n'
      'The rest of the pack contains only the new Day 5 learning material.\n\n'
      'TODAY\n'
      '1) Add 3 Power Query columns to FactSales: Gross Sales, Discount Amount, Discount Status.\n'
      '2) Close & Apply.\n'
      '3) Create basic DAX Measures.\n'
      '4) Build the first Sales Overview report.\n'
      '5) Add Slicers and test how the same Measures change with filters.\n\n'
      'FILES\n'
      '01_Nova_Day05_Start.pbix\n'
      '02_Power_Query_Calculations.xlsx\n'
      '03_First_DAX_Measures.xlsx\n'
      '04_Sales_Overview_Blueprint.xlsx\n'
      '05_Day05_Analyst_Challenge.xlsx\n\n'
      'SUCCESS = Clean model from Day 4 + new Power Query columns + Measures + first working report.\n',
      encoding='utf-8'
    )

    save_power_query(WORK/'02_Power_Query_Calculations.xlsx')
    save_dax(WORK/'03_First_DAX_Measures.xlsx')
    save_report(WORK/'04_Sales_Overview_Blueprint.xlsx')
    save_challenge(WORK/'05_Day05_Analyst_Challenge.xlsx')

    OUT.parent.mkdir(parents=True,exist_ok=True)
    if OUT.exists(): OUT.unlink()
    with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(WORK.rglob('*')):
            if p.is_file():
                z.write(p,p.relative_to(WORK))
    shutil.rmtree(WORK)
    print(f'Built {OUT}')

if __name__=='__main__':
    build()
