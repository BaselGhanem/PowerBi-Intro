from pathlib import Path
import base64, hashlib, shutil, zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'week2' / 'day6' / 'downloads' / 'Day06_CONTEXT_AND_CALCULATE.zip'
WORK = ROOT / 'week2' / 'day6' / '_generated_day6'
PBIX_PARTS = ROOT / 'tools' / 'day6_pbix_parts'
PBIX_SHA256 = '4068d4f4f38ae99f5b7924e26e8d968a2cdb14df4835cfbc44e7e6751e885f62'

TEAL='099999'
NAVY='13233B'
GOLD='F2C811'
WHITE='FFFFFF'
LIGHT='F3F7F7'
TEXT='1D2A33'
MUTED='61717C'
GREEN='1F8A70'
PURPLE='6D53C8'
PINK='D9476F'

thin = Side(style='thin', color='D8E2E5')

def title_block(ws, title, subtitle):
    ws.merge_cells('A1:F1')
    c=ws['A1']; c.value=title; c.font=Font(name='Almarai',size=20,bold=True,color=WHITE)
    c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='left',vertical='center')
    ws.row_dimensions[1].height=34
    ws.merge_cells('A2:F2')
    c=ws['A2']; c.value=subtitle; c.font=Font(name='Almarai',size=11,color=MUTED)
    c.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    ws.merge_cells('A3:F3')
    c=ws['A3']; c.value='X Academy • Power BI Specialist • Day 6'; c.font=Font(name='Almarai',size=10,bold=True,color=TEAL)

def add_table(ws, start_row, headers, rows, widths=None, row_height=40):
    for col,h in enumerate(headers,1):
        c=ws.cell(start_row,col,h)
        c.font=Font(name='Almarai',bold=True,color=WHITE)
        c.fill=PatternFill('solid',fgColor=TEAL)
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        c.border=Border(bottom=thin)
    for r_idx,row in enumerate(rows,start_row+1):
        fill=LIGHT if (r_idx-start_row)%2==1 else WHITE
        for col,val in enumerate(row,1):
            c=ws.cell(r_idx,col,val)
            c.font=Font(name='Almarai',size=10,color=TEXT)
            c.fill=PatternFill('solid',fgColor=fill)
            c.alignment=Alignment(vertical='top',wrap_text=True)
            c.border=Border(bottom=thin)
        ws.row_dimensions[r_idx].height=row_height
    if widths:
        for i,w in enumerate(widths,1):
            ws.column_dimensions[get_column_letter(i)].width=w

def new_wb(title, subtitle):
    wb=Workbook()
    ws=wb.active
    ws.title='Guide'
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='A5'
    title_block(ws,title,subtitle)
    return wb,ws

def restore_start_pbix(path):
    parts=sorted(PBIX_PARTS.glob('part*.b64'))
    if not parts:
        raise FileNotFoundError(f'No PBIX parts found in {PBIX_PARTS}')
    payload=''.join(p.read_text(encoding='utf-8').strip() for p in parts)
    raw=base64.b64decode(payload)
    sha=hashlib.sha256(raw).hexdigest()
    if sha != PBIX_SHA256:
        raise ValueError(f'PBIX checksum mismatch: {sha}')
    path.write_bytes(raw)

def save_filter_context(path):
    wb,ws=new_wb('FILTER CONTEXT LAB','نفس الـMeasure، نفس الـDAX، لكن النتيجة تتغير حسب الفلاتر المحيطة بها.')
    rows=[
      ['01','No selection','[Total Revenue]','كل البيانات الحالية','هذه هي نقطة البداية.'],
      ['02','Month slicer → March','[Total Revenue]','Revenue لشهر March فقط','هل غيرنا صيغة الـMeasure؟'],
      ['03','Region slicer → Amman','[Total Revenue]','Revenue لمنطقة Amman فقط','ما الـFilter الموجود حاليا؟'],
      ['04','Click a Category bar','[Total Revenue]','Revenue للفئة المختارة','الـVisual نفسه يستطيع إنشاء Context.'],
      ['05','Month + Region + Category','[Total Revenue]','تقاطع كل الفلاتر معا','Measure + Current Filters = Result']
    ]
    add_table(ws,5,['Step','Action','Same Measure','Result means','Ask the learner'],rows,[8,28,24,42,46],44)

    ref=wb.create_sheet('Day5 Reference')
    ref.sheet_view.showGridLines=False
    title_block(ref,'WHAT WE ALREADY HAVE','هذا الملف يبدأ من نهاية Day 5 كما تم بناؤها في الصف.')
    pq=[
      ['Power Query','Discount Amount','[Discount %] * [Revenue]'],
      ['Power Query','Sales Value','[Unit Price] * [Quantity]'],
      ['Power Query','Discount Status','if [Discount Amount] = 0 then "Full Price" else "Discounted Price"'],
      ['Measure','Total Revenue','Existing Day 5 measure'],
      ['Measure','Transactions','Existing Day 5 measure'],
      ['Measure','Customer Count','Existing Day 5 measure'],
      ['Measure','Total QTY','Existing Day 5 measure']
    ]
    add_table(ref,5,['Type','Name','Logic / Status'],pq,[20,28,80],42)
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def save_calculate(path):
    wb,ws=new_wb('CALCULATE MEASURES','CALCULATE لا تعني معادلة جديدة فقط؛ هي تعيد تقييم Measure تحت Filter Context معدل.')
    rows=[
      ['01','Discounted Revenue','CALCULATE','Discounted Revenue =\nCALCULATE(\n    [Total Revenue],\n    FactSales[Discount Status] = "Discounted Price"\n)','Add filter','Revenue من الصفوف Discounted Price فقط','Required'],
      ['02','Full Price Revenue','CALCULATE','Full Price Revenue =\nCALCULATE(\n    [Total Revenue],\n    FactSales[Discount Status] = "Full Price"\n)','Add filter','Revenue من Full Price فقط','Required'],
      ['03','Retail Revenue','CALCULATE','Retail Revenue =\nCALCULATE(\n    [Total Revenue],\n    DimBranch[Channel] = "Retail"\n)','Add filter','Revenue لقناة Retail مع احترام Month/Region/Category','Required'],
      ['04','Revenue All Categories','CALCULATE + REMOVEFILTERS','Revenue All Categories =\nCALCULATE(\n    [Total Revenue],\n    REMOVEFILTERS(DimProduct[Category])\n)','Remove one filter','يلغي Category فقط ويبقي بقية Context','Required'],
      ['05','Category Contribution %','DIVIDE','Category Contribution % =\nDIVIDE(\n    [Total Revenue],\n    [Revenue All Categories]\n)','Reuse measures','حصة كل Category من Revenue داخل الفلاتر الحالية','Required'],
      ['06','Discounted Transactions','CALCULATE','Discounted Transactions =\nCALCULATE(\n    [Transactions],\n    FactSales[Discount Status] = "Discounted Price"\n)','Add filter','عدد Transactions المخفضة','Optional']
    ]
    add_table(ws,5,['#','Measure','Function','DAX','Context action','Business meaning','Level'],rows,[6,28,25,58,24,46,14],78)
    ws['A13']='احفظها بهذه الطريقة'
    ws['A13'].font=Font(name='Almarai',size=12,bold=True,color=TEAL)
    ws.merge_cells('A14:F15')
    ws['A14']='CALCULATE = خذ الـMeasure الموجودة، وعدل الـFilter Context، ثم احسبها من جديد. REMOVEFILTERS لا يمسح كل شيء؛ هنا نزيل Category فقط ونبقي Month وRegion وBranch وغيرها.'
    ws['A14'].font=Font(name='Almarai',size=12,bold=True,color=TEXT)
    ws['A14'].fill=PatternFill('solid',fgColor='E7F5F4')
    ws['A14'].alignment=Alignment(wrap_text=True,vertical='center')
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def save_blueprint(path):
    wb,ws=new_wb('CONTEXT LAB — REPORT BLUEPRINT','ابن صفحة ثانية فوق تقرير Day 5 لتشاهد كيف تتغير الـMeasures وكيف تتفاعل الـVisuals.')
    layout=[
      ['Page','Context Lab','','','Duplicate Day 5 report page or start a clean page'],
      ['Card 1','Card','[Total Revenue]','','Baseline measure'],
      ['Card 2','Card','[Discounted Revenue]','','CALCULATE adds Discount Status filter'],
      ['Card 3','Card','[Full Price Revenue]','','Compare with discounted revenue'],
      ['Card 4','Card','[Retail Revenue]','','CALCULATE using DimBranch[Channel]'],
      ['Visual 1','Bar Chart','DimBranch[Region]','[Total Revenue]','Click a Region and watch the context change'],
      ['Visual 2','Matrix','DimProduct[Category]','[Total Revenue] + [Category Contribution %]','Contribution should make business sense'],
      ['Visual 3','Line Chart','DimDate[Month]','[Total Revenue]','Keep Month sorted by Month Number'],
      ['Slicer 1','Slicer','DimDate[Month]','',''],
      ['Slicer 2','Slicer','DimBranch[Region]','',''],
      ['Slicer 3','Slicer','DimProduct[Category]','','']
    ]
    add_table(ws,5,['Element','Visual','Axis / Field','Values','Teaching point'],layout,[14,22,34,40,58],44)

    inter=wb.create_sheet('Interactions Lab')
    inter.sheet_view.showGridLines=False
    title_block(inter,'VISUAL INTERACTIONS','الـVisual ليس مجرد صورة؛ يمكن أن يغير Context لبقية الصفحة.')
    rows=[
      ['1','Click one Category','Observe Cards + Region + Trend','Cross-filter / highlight happens automatically'],
      ['2','Click one Region','Observe Category + Cards','Explain that a chart selection is another filter source'],
      ['3','Format → Edit interactions','Choose one target visual','Switch between Filter / Highlight / None'],
      ['4','Set one interaction to None','Explain why you chose it','Intentional design, not random formatting'],
      ['5','Reset selections','Return to All','Always verify the full-page state']
    ]
    add_table(inter,5,['Step','Action','Observe','Meaning'],rows,[8,36,44,70],44)
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def save_challenge(path):
    wb,ws=new_wb('DAY 6 — CONTEXT CHALLENGE','ابن Context Analysis واثبت أنك تفهم لماذا تغيرت النتيجة، لا فقط كيف تكتب DAX.')
    req=[
      ['01','Create required Measures','Discounted Revenue • Full Price Revenue • Retail Revenue • Revenue All Categories • Category Contribution %'],
      ['02','Build 4 Cards','Total Revenue • Discounted Revenue • Full Price Revenue • Retail Revenue'],
      ['03','Build Category Matrix','Category + Total Revenue + Category Contribution %'],
      ['04','Build Region Analysis','Revenue by Region'],
      ['05','Use 3 Slicers','Month • Region • Category'],
      ['06','Change one interaction','Use Edit interactions and intentionally set one target to None'],
      ['07','Explain the result','Do not say "Power BI changed it"; name the Filter Context']
    ]
    add_table(ws,5,['#','Requirement','Proof'],req,[8,32,95],46)
    ws['A14']='Business Questions'
    ws['A14'].font=Font(name='Almarai',size=12,bold=True,color=TEAL)
    q=[
      ['1','Why does [Total Revenue] change when Region changes although the DAX formula stays the same?'],
      ['2','What exactly did CALCULATE change in [Discounted Revenue]?'],
      ['3','Why does [Revenue All Categories] still react to Month and Region?'],
      ['4','Select one Region and one Month. Do Category Contribution percentages still describe the share inside that selection?'],
      ['5','What is the difference between a slicer filter and clicking a chart from the Measure point of view?'],
      ['6','Why would you intentionally set one visual interaction to None?']
    ]
    add_table(ws,15,['#','Question'],q,[8,110],48)
    ws.merge_cells('A23:F25')
    ws['A23']='Bridge to Day 7: الآن فهمنا أن الزمن نفسه Filter Context. في Day 7 سنستخدم DimDate وTime Intelligence لمقارنة Current vs Previous Period وYTD.'
    ws['A23'].font=Font(name='Almarai',size=12,bold=True,color=WHITE)
    ws['A23'].fill=PatternFill('solid',fgColor=TEAL)
    ws['A23'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def build():
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)

    restore_start_pbix(WORK/'01_Nova_Day06_Start.pbix')
    save_filter_context(WORK/'02_Filter_Context_Lab.xlsx')
    save_calculate(WORK/'03_CALCULATE_Measures.xlsx')
    save_blueprint(WORK/'04_Context_Report_Blueprint.xlsx')
    save_challenge(WORK/'05_Day06_Analyst_Challenge.xlsx')

    (WORK/'00_START_HERE.txt').write_text(
      'NOVA DISTRIBUTION GROUP — DAY 06: CONTEXT CHANGES EVERYTHING\n\n'
      'START\n'
      'Open 01_Nova_Day06_Start.pbix. It is the completed Day 5 file supplied by the trainer and is our shared starting point.\n'
      'No new dataset is introduced today.\n\n'
      'DAY 6 STORY\n'
      'Day 5 proved that the same Measure changes when we use slicers and click visuals. Today we explain WHY.\n\n'
      'FLOW — ABOUT 2 HOURS\n'
      '00–10 min  Review Day 5 report and existing Measures.\n'
      '10–30 min  Filter Context: same Measure + different filters = different result.\n'
      '30–55 min  CALCULATE: add a filter using Discount Status and Channel.\n'
      '55–75 min  REMOVEFILTERS + Category Contribution %.\n'
      '75–100 min Build Context Lab page.\n'
      '100–110 min Cross-filtering + Edit interactions.\n'
      '110–120 min Analyst Challenge + bridge to Time Intelligence.\n\n'
      'TODAY WE DO NOT TEACH\n'
      'FILTER table function, VAR, SELECTEDVALUE, ALLSELECTED, Dynamic Titles or advanced time intelligence. Keep the mental model clean.\n\n'
      'SUCCESS\n'
      'The learner can explain why a Measure changed, use CALCULATE to add a filter, use REMOVEFILTERS to remove one filter, calculate Category Contribution %, and intentionally control a visual interaction.\n',
      encoding='utf-8'
    )

    OUT.parent.mkdir(parents=True,exist_ok=True)
    if OUT.exists():
        OUT.unlink()
    with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(WORK.rglob('*')):
            if p.is_file():
                z.write(p,p.relative_to(WORK))
    shutil.rmtree(WORK)
    print(f'Built {OUT}')

if __name__=='__main__':
    build()
