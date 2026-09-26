from pathlib import Path
from datetime import date, timedelta
import sys, zipfile, shutil

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

TOOLS = Path(__file__).resolve().parent
ROOT = TOOLS.parent
sys.path.insert(0, str(TOOLS))
import generate_week1_data as w1

OUT = ROOT / 'week2' / 'day4' / 'downloads' / 'Day04_MODEL_THE_BUSINESS.zip'
WORK = ROOT / 'week2' / 'day4' / '_generated_day4'

SALES_HEADER = ['Transaction ID','Date','Customer ID','Product ID','Sales Rep ID','Branch ID','Region','Channel','Quantity','Unit Price','Discount %','Revenue']

def write_xlsx(path, header, rows, sheet='Data'):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = sheet[:31]
    header_fill = PatternFill('solid', fgColor='13233B')
    for c, value in enumerate(header, 1):
        cell = ws.cell(1, c, value)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center')
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        max_len = min(max(len(str(c.value or '')) for c in list(col)[:250]) + 2, 28)
        ws.column_dimensions[letter].width = max(12, max_len)
    wb.save(path)

def add_sheet(wb, title, header, rows):
    ws = wb.create_sheet(title[:31])
    fill = PatternFill('solid', fgColor='13233B')
    for c, value in enumerate(header, 1):
        cell = ws.cell(1, c, value)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws

def model_workbook(path, fact, products, customers, branches):
    wb = Workbook()
    wb.remove(wb.active)
    fact_header = ['Transaction ID','Date','Customer ID','Product ID','Branch ID','Quantity','Unit Price','Discount %','Revenue']
    add_sheet(wb, 'FactSales', fact_header, fact)

    product_rows = [[p['product_id'],p['product_name'],p['category'],p['subcategory'],p['brand'],p['unit_cost'],p['list_price'],p['status']] for p in products]
    add_sheet(wb, 'DimProduct', ['Product ID','Product Name','Category','Subcategory','Brand','Unit Cost','List Price','Status'], product_rows)

    customer_rows = [[c['customer_id'],c['customer_name'],c['city'],c['region'],c['segment'],c['customer_type'],c['status']] for c in customers]
    add_sheet(wb, 'DimCustomer', ['Customer ID','Customer Name','City','Region','Segment','Customer Type','Status'], customer_rows)

    branch_rows = [[b['branch_id'],b['branch_name'],b['city'],b['region'],b['channel'],b['manager'],b['status']] for b in branches]
    add_sheet(wb, 'DimBranch', ['Branch ID','Branch Name','City','Region','Channel','Branch Manager','Status'], branch_rows)

    start=date(2026,1,1); end=date(2026,4,30)
    names=['January','February','March','April','May','June','July','August','September','October','November','December']
    d=[]; cur=start
    while cur<=end:
        d.append([cur.isoformat(),cur.year,cur.month,names[cur.month-1],f'Q{((cur.month-1)//3)+1}',f'{cur.year}-{cur.month:02d}'])
        cur += timedelta(days=1)
    add_sheet(wb, 'DimDate', ['Date','Year','Month Number','Month','Quarter','Year Month'], d)

    for ws in wb.worksheets:
        for col in ws.columns:
            letter=col[0].column_letter
            max_len=min(max(len(str(c.value or '')) for c in list(col)[:250])+2,28)
            ws.column_dimensions[letter].width=max(12,max_len)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

def build():
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)

    products=w1.product_catalog_rows()
    customers=w1.customer_catalog_rows()
    branches=w1.branch_catalog_rows()

    months={m:list(w1.monthly_sales_rows(10000,m,products,customers,branches)) for m in range(1,6)}

    # 01 — Merge review
    merge_dir=WORK/'01_REVIEW_MERGE'
    write_xlsx(merge_dir/'01_Sales_Sample.xlsx', SALES_HEADER, months[1][:1200], 'Sales')
    product_header=['Product ID','SKU','Product Name','Category','Subcategory','Brand','Pack / Variant','Unit Cost','List Price','Supplier','Portfolio Role','Status']
    product_rows=[[p['product_id'],p['sku'],p['product_name'],p['category'],p['subcategory'],p['brand'],p['variant'],p['unit_cost'],p['list_price'],p['supplier'],p['portfolio_role'],p['status']] for p in products]
    write_xlsx(merge_dir/'02_Product_Master.xlsx', product_header, product_rows, 'Products')

    # 02 — Append review
    append_dir=WORK/'02_REVIEW_APPEND'
    for m,name in [(1,'January'),(2,'February'),(3,'March')]:
        write_xlsx(append_dir/f'Sales_{name}.xlsx', SALES_HEADER, months[m][:2500], 'Sales')

    # 03 — Folder review with Excel and May held out
    folder_dir=WORK/'03_REVIEW_FOLDER'/'Monthly_Sales_Excel'
    for m,name in [(1,'Jan'),(2,'Feb'),(3,'Mar'),(4,'Apr')]:
        write_xlsx(folder_dir/f'Sales_{name}.xlsx', SALES_HEADER, months[m], 'Sales')
    new_dir=WORK/'03_REVIEW_FOLDER'/'NEW_MONTH_TO_DROP_AFTER_BUILD'
    write_xlsx(new_dir/'Sales_May.xlsx', SALES_HEADER, months[5], 'Sales')

    # 04 — Model lab. Same Jan-Apr business universe, shaped for Star Schema.
    fact=[]
    for m in range(1,5):
        for r in months[m]:
            fact.append([r[0],r[1],r[2],r[3],r[5],r[8],r[9],r[10],r[11]])
    model_workbook(WORK/'04_MODEL_LAB'/'Nova_Day04_Model_Lab.xlsx', fact, products, customers, branches)

    # Checklist/reference
    w1.make_simple_xlsx(WORK/'05_Day04_Model_Checklist.xlsx', [
        ('Model Steps', [
            ['Step','What to do','Proof'],
            ['1','Identify FactSales grain','One row = one transaction'],
            ['2','Check dimension keys','Product ID / Customer ID / Branch ID / Date are unique in their dimensions'],
            ['3','Create relationships','Each Dimension 1 → * FactSales'],
            ['4','Set filter direction','Single: Dimension → Fact'],
            ['5','Arrange Star Schema','FactSales center, Dimensions around it'],
            ['6','Test filters','Slicers from each Dimension change Sales results'],
            ['7','Finish','Model is Ready for DAX']
        ]),
        ('Relationship Map', [
            ['Dimension','Key (1 side)','Fact column (* side)','Cardinality','Filter Direction'],
            ['DimProduct','Product ID','FactSales[Product ID]','1:*','Single → FactSales'],
            ['DimCustomer','Customer ID','FactSales[Customer ID]','1:*','Single → FactSales'],
            ['DimBranch','Branch ID','FactSales[Branch ID]','1:*','Single → FactSales'],
            ['DimDate','Date','FactSales[Date]','1:*','Single → FactSales']
        ]),
        ('Memory', [
            ['Concept','Remember'],
            ['Merge','Add Columns'],
            ['Append','Add Rows'],
            ['Folder','Repeat the same process for new files'],
            ['Fact','What happened?'],
            ['Dimension','Describe what happened'],
            ['Star Schema','Fact in center, Dimensions around it']
        ])
    ])

    # QA before packing
    assert len(fact)==40000
    assert len({p['product_id'] for p in products})==len(products)==240
    assert len({c['customer_id'] for c in customers})==len(customers)==8000
    assert len({b['branch_id'] for b in branches})==len(branches)==100
    assert all(r[3] in {p['product_id'] for p in products} for r in fact)
    assert all(r[2] in {c['customer_id'] for c in customers} for r in fact)
    assert all(r[4] in {b['branch_id'] for b in branches} for r in fact)

    (WORK/'00_START_HERE.txt').write_text(
        'NOVA DISTRIBUTION GROUP — DAY 04: MODEL THE BUSINESS\n\n'
        'STORY\n'
        'Week 1 gave us connected, clean and refreshable data. Today we first review Merge, Append and Folder using Excel files, then build a real Star Schema.\n\n'
        'PART A — REVIEW\n'
        '01_REVIEW_MERGE: Merge Sales with Product Master on Product ID. Remember: Merge = Add Columns.\n'
        '02_REVIEW_APPEND: Append January + February + March. Remember: Append = Add Rows.\n'
        '03_REVIEW_FOLDER: Connect to Monthly_Sales_Excel as Folder, combine Jan-Apr, then move Sales_May.xlsx into the folder and Refresh.\n\n'
        'PART B — MODEL\n'
        'Open 04_MODEL_LAB/Nova_Day04_Model_Lab.xlsx and load FactSales, DimProduct, DimCustomer, DimBranch and DimDate.\n'
        'Build 1:* relationships from each Dimension to FactSales. Use single filter direction. Arrange a Star Schema and test filters.\n\n'
        'SUCCESS = FactSales + 4 Dimensions + working filters + READY FOR DAX.\n',
        encoding='utf-8'
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    if OUT.exists(): OUT.unlink()
    with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(WORK.rglob('*')):
            if p.is_file():
                z.write(p,p.relative_to(WORK))
    shutil.rmtree(WORK)
    print(f'Built {OUT}')

if __name__=='__main__':
    build()
