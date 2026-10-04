"""Build the Day 7 four-role lab. Preserve original Day 4 data byte-for-byte."""
from pathlib import Path
from datetime import date,timedelta
import csv,json,io,random,calendar,zipfile,base64,zlib,hashlib,html,shutil,uuid
from openpyxl import load_workbook
from PIL import Image
ROOT=Path(__file__).resolve().parent.parent
WORK=ROOT/'week3/day7/_generated_day7'
OUT=ROOT/'week3/day7/downloads/Day07_FOUR_DASHBOARDS.zip'
SCHEMA='https://developer.microsoft.com/json-schemas/'
RPT=SCHEMA+'fabric/item/report/definition/'
RNG=random.Random(72026)
DATA={};TYPES={};MEAS={}
def write_json(p,o):
 p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(o,ensure_ascii=False,indent=2),encoding='utf-8')
def table(name,headers,rows,types):
 DATA[name]=(headers,list(rows));TYPES[name]=types
 p=WORK/'01_DATA'/f'{name}.csv';p.parent.mkdir(parents=True,exist_ok=True)
 with p.open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.writer(f);w.writerow(headers);w.writerows(DATA[name][1])
def measure(table,name,formula,fmt='#,##0',level='base',meaning=''):
 MEAS.setdefault(table,[]).append(dict(name=name,expression=formula,formatString=fmt,displayFolder='Day 7/'+('01 Base' if level=='base' else '02 Time'),level=level,meaning=meaning))
def build_data():
 with zipfile.ZipFile(ROOT/'week2/day4/downloads/Day04_MODEL_THE_BUSINESS.zip') as z:
  raw=z.read('04_MODEL_LAB/Nova_Day04_Model_Lab.xlsx')
 p=WORK/'00_DAY04_ORIGINAL/Nova_Day04_Model_Lab.xlsx';p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(raw)
 wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 for n in ['DimProduct','DimCustomer','DimBranch']:
  it=wb[n].values;header=list(next(it));rows=list(it)
  types=['double' if isinstance(rows[0][i],(float,int)) else 'string' for i in range(len(header))]
  table(n,header,rows,types)
 it=wb['FactSales'].values;header=list(next(it));original=[list(r) for r in it]
 header+=['Gross Sales','Discount Amount','Discount Status','Source']
 rows=[]
 for r in original:
  gross=round(r[5]*r[6],2);disc=round(gross-r[8],2)
  rows.append(r+[gross,disc,'Full Price' if disc==0 else 'Discounted Price','Day04 original'])
 # Comparable 2025 has the same transaction count, customer/product/branch mapping.
 # Quantity and discount vary by month and category. This is explicitly synthetic.
 cat={r[0]:r[2] for r in DATA['DimProduct'][1]}
 for r in original:
  rr=r.copy();rr[0]='PY-'+str(r[0]);rr[1]='2025'+str(r[1])[4:]
  month=int(str(rr[1])[5:7]);factor=[.89,1.02,.93,1.08][month-1]
  # Category-dependent movement makes growth analysis more than a universal uplift.
  factor*=1+((sum(ord(c) for c in cat[rr[3]])%7)-3)*.025
  rr[5]=max(1,round(rr[5]*factor));rr[6]=round(rr[6]*.96,2)
  gross=round(rr[5]*rr[6],2);rr[8]=round(gross*(1-rr[7]),2);disc=round(gross-rr[8],2)
  rows.append(rr+[gross,disc,'Full Price' if disc==0 else 'Discounted Price','Synthetic 2025 comparison'])
 table('FactSales',header,rows,['string','dateTime','string','string','string','int64','double','double','double','double','double','string','string'])
 dates=[];d=date(2025,1,1)
 while d<=date(2026,12,31):
  dates.append([d.isoformat(),d.year,d.month,calendar.month_name[d.month],f'Q{(d.month-1)//3+1}',d.strftime('%Y-%m'),d.year*100+d.month]);d+=timedelta(days=1)
 table('DimDate',['Date','Year','Month Number','Month','Quarter','Year Month','Year Month Sort'],dates,['dateTime','int64','int64','string','string','string','int64'])
 departments=['Sales','Operations','Finance','HR','Planning']
 table('DimDepartment',['Department ID','Department'],[[f'D{i+1}',n] for i,n in enumerate(departments)],['string','string'])
 branches=DATA['DimBranch'][1];products=DATA['DimProduct'][1]
 ops=[];fin=[];work=[]
 for year in [2025,2026]:
  for month in range(1,5):
   end=calendar.monthrange(year,month)[1]
   for i in range(1200):
    dt=date(year,month,RNG.randint(1,end));branch=branches[RNG.randrange(len(branches))][0]
    state=RNG.choices(['Completed','Open','Cancelled'],weights=[85,11,4])[0]
    duration=RNG.randint(1,7);sla=3
    ops.append([f'O{year}{month:02}{i:05}',dt.isoformat(),branch,state,duration,sla,int(state=='Completed' and duration<=sla),RNG.randint(10,250),RNG.choice(['Standard','Express']),f'OP{RNG.randint(1,8):02}'])
   for branch in branches:
    dt=date(year,month,end).isoformat();b=branch[0]
    income=round(RNG.uniform(25000,65000)*(1+.035*month)* (1.08 if year==2026 else 1),2)
    cost=round(income*RNG.uniform(.65,.91),2);budget=round(income*RNG.uniform(.94,1.13),2)
    fin.append([f'F{year}{month:02}{b}',dt,b,income,cost,budget])
    for j,dep in enumerate(departments):
     planned=RNG.randint(10,34);actual=max(1,planned-RNG.randint(-2,6));capacity=actual*22*8
     demand=round(planned*22*8*RNG.uniform(.73,1.04),1)
     work.append([dt,b,f'D{j+1}',planned,actual,capacity,demand,RNG.randint(0,actual*3)])
 table('FactOperations',['Order ID','Order Date','Branch ID','Status','Cycle Days','SLA Days','On Time Flag','Units','Priority','Operator ID'],ops,['string','dateTime','string','string','int64','int64','int64','int64','string','string'])
 table('FactFinance',['Record ID','Posting Date','Branch ID','Revenue','Cost','Revenue Budget'],fin,['string','dateTime','string','double','double','double'])
 table('FactWorkforce',['Snapshot Date','Branch ID','Department ID','Planned Employees','Actual Headcount','Capacity Hours','Demand Hours','Absence Days'],work,['dateTime','string','string','int64','int64','double','double','int64'])
 return raw

def build_measures():
 measure('FactSales','Total Revenue','SUM(FactSales[Revenue])','#,##0.00',meaning='Net revenue after discount in JOD.')
 measure('FactSales','Transactions','COUNTROWS(FactSales)',meaning='One row is one sales transaction.')
 measure('FactSales','Customer Count','DISTINCTCOUNT(FactSales[Customer ID])')
 measure('FactSales','Total Quantity','SUM(FactSales[Quantity])')
 measure('FactSales','Average Transaction Value','DIVIDE([Total Revenue], [Transactions])','#,##0.00')
 measure('FactOperations','Orders','COUNTROWS(FactOperations)')
 measure('FactOperations','Completed Orders','CALCULATE([Orders], FactOperations[Status] = "Completed")')
 measure('FactOperations','On Time Orders','CALCULATE([Orders], FactOperations[Status] = "Completed", FactOperations[On Time Flag] = 1)')
 measure('FactOperations','On Time %','DIVIDE([On Time Orders], [Completed Orders])','0.0%',meaning='Only completed orders are in the denominator; cancelled and open excluded.')
 measure('FactOperations','Average Cycle Days','CALCULATE(AVERAGE(FactOperations[Cycle Days]), FactOperations[Status] = "Completed")','0.0',meaning='Completed orders grouped by order date, not completion date.')
 measure('FactOperations','Open Orders','CALCULATE([Orders], FactOperations[Status] = "Open")')
 for name,col in [('Finance Revenue','Revenue'),('Total Cost','Cost'),('Total Revenue Budget','Revenue Budget')]:measure('FactFinance',name,f'SUM(FactFinance[{col}])','#,##0.00')
 measure('FactFinance','Operating Profit','[Finance Revenue] - [Total Cost]','#,##0.00')
 measure('FactFinance','Operating Margin %','DIVIDE([Operating Profit], [Finance Revenue])','0.0%')
 measure('FactFinance','Budget Variance','[Finance Revenue] - [Total Revenue Budget]','#,##0.00',meaning='Positive is favorable for revenue.')
 measure('FactFinance','Budget Achievement %','DIVIDE([Finance Revenue], [Total Revenue Budget])','0.0%')
 measure('FactWorkforce','Last Snapshot Date','MAX(FactWorkforce[Snapshot Date])','yyyy-MM-dd',meaning='Last available snapshot within the current date, branch and department context.')
 for name,col in [('Headcount','Actual Headcount'),('Planned Headcount','Planned Employees'),('Available Hours','Capacity Hours'),('Required Hours','Demand Hours'),('Absent Days','Absence Days')]:
  measure('FactWorkforce',name,f'VAR LastSnapshot = [Last Snapshot Date]\nRETURN\n    IF(NOT ISBLANK(LastSnapshot), CALCULATE(SUM(FactWorkforce[{col}]), KEEPFILTERS(FactWorkforce[Snapshot Date] = LastSnapshot)))',meaning='Last monthly snapshot only; never summed across months.')
 measure('FactWorkforce','Staffing Gap','[Headcount] - [Planned Headcount]',meaning='Negative means understaffed.')
 measure('FactWorkforce','Staffing Coverage %','DIVIDE([Headcount], [Planned Headcount])','0.0%')
 measure('FactWorkforce','Capacity Coverage %','DIVIDE([Available Hours], [Required Hours])','0.0%')
 # A shared pattern across roles. HR is a snapshot: no YTD sum of headcount.
 for tab,base,prefix in [('FactSales','Total Revenue','Revenue'),('FactOperations','Completed Orders','Completed Orders'),('FactFinance','Finance Revenue','Finance Revenue'),('FactWorkforce','Headcount','Headcount')]:
  fmt='#,##0.00' if tab in ['FactSales','FactFinance'] else '#,##0'
  measure(tab,prefix+' PM',f'IF(HASONEVALUE(DimDate[Year Month]), CALCULATE([{base}], DATEADD(DimDate[Date], -1, MONTH)))',fmt,'time','Requires one month selected; first data month returns blank.')
  measure(tab,prefix+' MoM %',f'DIVIDE([{base}] - [{prefix} PM], [{prefix} PM])','0.0%','time')
  measure(tab,prefix+' PY',f'CALCULATE([{base}], SAMEPERIODLASTYEAR(DimDate[Date]))',fmt,'time')
  measure(tab,prefix+' YoY %',f'DIVIDE([{base}] - [{prefix} PY], [{prefix} PY])','0.0%','time')
  if tab!='FactWorkforce':
   measure(tab,prefix+' YTD',f'IF(HASONEVALUE(DimDate[Year]), CALCULATE([{base}], DATESYTD(DimDate[Date])))',fmt,'time','Requires one year; select a month to set the cutoff.')
   measure(tab,prefix+' PY YTD',f'IF(HASONEVALUE(DimDate[Year]), CALCULATE([{prefix} YTD], SAMEPERIODLASTYEAR(DimDate[Date])))',fmt,'time')
   measure(tab,prefix+' YTD Growth %',f'DIVIDE([{prefix} YTD] - [{prefix} PY YTD], [{prefix} PY YTD])','0.0%','time')
 measure('FactOperations','On Time % PM','IF(HASONEVALUE(DimDate[Year Month]), CALCULATE([On Time %], DATEADD(DimDate[Date], -1, MONTH)))','0.0%','time')
 measure('FactOperations','On Time Change pp','IF(NOT ISBLANK([On Time % PM]), ([On Time %] - [On Time % PM]) * 100)','0.0 "pp"','time','Percentage-point difference, not relative growth.')
 measure('FactWorkforce','Staffing Coverage % PM','IF(HASONEVALUE(DimDate[Year Month]), CALCULATE([Staffing Coverage %], DATEADD(DimDate[Date], -1, MONTH)))','0.0%','time')
 measure('FactWorkforce','Staffing Coverage Change pp','IF(NOT ISBLANK([Staffing Coverage % PM]), ([Staffing Coverage %] - [Staffing Coverage % PM]) * 100)','0.0 "pp"','time')

ROLES=[
 dict(name='MAHMOUD',ar='محمود',domain='SALES',table='FactSales',base='Total Revenue',prefix='Revenue',cards=['Total Revenue','Revenue PY','Revenue YoY %','Customer Count'],axis=('DimProduct','Category'),bar='Total Revenue',question='أي فئة ارتفعت مبيعاتها مقارنة بالسنة الماضية؟',slicer=('DimProduct','Category')),
 dict(name='SOMAYA',ar='سمية',domain='OPERATIONS',table='FactOperations',base='Completed Orders',prefix='Completed Orders',cards=['Orders','Completed Orders','On Time %','Average Cycle Days'],axis=('DimBranch','Region'),bar='Completed Orders',question='أي منطقة تحتاج تحسين الالتزام بوقت التسليم؟',slicer=('FactOperations','Priority')),
 dict(name='HISHAM',ar='هشام',domain='FINANCE',table='FactFinance',base='Finance Revenue',prefix='Finance Revenue',cards=['Finance Revenue','Operating Profit','Operating Margin %','Budget Achievement %'],axis=('DimBranch','Region'),bar='Budget Variance',question='أين الإيرادات أقل من الموازنة رغم نموها؟',slicer=('DimBranch','Channel')),
 dict(name='ENAS',ar='إيناس',domain='WORKFORCE',table='FactWorkforce',base='Headcount',prefix='Headcount',cards=['Headcount','Planned Headcount','Staffing Gap','Staffing Coverage %'],axis=('DimDepartment','Department'),bar='Staffing Gap',question='أي قسم يحتاج زيادة الموارد لتغطية العمل؟',slicer=('DimDepartment','Department'))
]

def expression(tab,prop,kind='Column'):return {kind:{'Expression':{'SourceRef':{'Entity':tab}},'Property':prop}}
def projection(tab,prop,kind='Column'):
 return {'field':expression(tab,prop,kind),'queryRef':tab+'.'+prop,'nativeQueryRef':prop,'active':True}
def lit(s):return {'expr':{'Literal':{'Value':s}}}
def color(s):return {'solid':{'color':lit("'"+s+"'")}}
def textbox(lines,x,y,w,h,size=16):
 return dict(visualType='textbox',objects={'general':[{'properties':{'paragraphs':[{'textRuns':[{'value':v,'textStyle':{'fontFamily':'Segoe UI','fontSize':f'{size}pt','color':'#13233B'}}]} for v in lines]}}]})
def chart(vtype,roles,title):
 return {'visualType':vtype,'query':{'queryState':{role:{'projections':[projection(*f) for f in fields]} for role,fields in roles.items()}},'visualContainerObjects':{'title':[{'properties':{'show':lit('true'),'text':lit("'"+title.replace("'","''")+"'"),'fontSize':lit('12D')}}],'background':[{'properties':{'show':lit('true'),'color':color('#FFFFFF'),'transparency':lit('0D')}}]},'drillFilterOtherVisuals':True}
def add_visual(pg,vis,x,y,w,h,i):
 name='v'+str(i).zfill(4)
 write_json(pg/'visuals'/name/'visual.json',{'$schema':RPT+'visualContainer/2.9.0/schema.json','name':name,'position':dict(x=x,y=y,width=w,height=h,z=i,tabOrder=i),'visual':vis})
def build_report(folder,stage):
 report=folder/'Nova.Report';definition=report/'definition'
 write_json(report/'definition.pbir',{'$schema':SCHEMA+'fabric/item/report/definitionProperties/2.0.0/schema.json','version':'4.0','datasetReference':{'byPath':{'path':'../Nova.SemanticModel'}}})
 write_json(definition/'version.json',{'$schema':RPT+'versionMetadata/1.0.0/schema.json','version':'2.0.0'})
 write_json(definition/'report.json',{'$schema':RPT+'report/3.3.0/schema.json','themeCollection':{}})
 order=['start','demo']+[r['name'].lower() for r in ROLES]
 write_json(definition/'pages/pages.json',{'$schema':RPT+'pagesMetadata/1.1.0/schema.json','pageOrder':order,'activePageName':'start'})
 for idx,name in enumerate(order):
  pg=definition/'pages'/name
  display='00 START HERE' if idx==0 else ('01 BASEL DEMO' if idx==1 else f'{idx:02} {ROLES[idx-2]["name"]} | {ROLES[idx-2]["domain"]}')
  write_json(pg/'page.json',{'$schema':RPT+'page/2.1.0/schema.json','name':name,'displayName':display,'displayOption':'FitToPage','width':1280,'height':720,'objects':{'background':[{'properties':{'color':color('#F3F7F7'),'transparency':lit('0D')}}]}})
  add_visual(pg,textbox([display,'NOVA  /  X Academy  /  Day 7'],24,12,1232,74,22),24,12,1232,74,0)
  if idx==0:
   lines=['1. Refresh to load the embedded data. No file-path setup is needed.','2. Mark DimDate as Date table using Date, then open BASEL DEMO. Select Year 2026.','3. Open your named page and the matching learner guide.','4. Build cards, trend, bar chart and matrix in the provided spaces.','5. Use DimDate fields for all time slicers. Sort Month by Month Number.','6. Checkpoints: 15 min cards / 30 min comparison / 45 min trend & bars.','Data: Jan-Apr 2025 and Jan-Apr 2026. Training data, not company records.','Original Day 4 sales are preserved. 2025 comparison and other domains are synthetic.','Workforce is a monthly snapshot: do not sum headcount across months.','Desktop opening and refresh must be checked by the trainer before class.']
   add_visual(pg,textbox(lines,32,110,1210,560,18),32,110,1210,560,1);continue
  role=ROLES[0] if idx==1 else ROLES[idx-2]
  tab=role['table'];base=role['base'];prefix=role['prefix'];comparison='Planned Headcount' if tab=='FactWorkforce' else prefix+' PY'
  for j,(t,c) in enumerate([('DimDate','Year'),('DimDate','Year Month'),role['slicer']]):
   add_visual(pg,chart('slicer',{'Values':[(t,c,'Column')]},c),24+j*414,92,400,80,j+1)
  if stage=='START' and idx!=1:
   for j,card in enumerate(role['cards']):add_visual(pg,textbox(['KPI '+str(j+1),card],0,0,0,0,16),24+j*310,192,296,112,j+4)
   add_visual(pg,textbox(['MONTHLY TREND',f'DimDate[Year Month] + [{base}] + [{comparison}]'],0,0,0,0),24,326,740,230,8)
   add_visual(pg,textbox(['BAR CHART',f'{role["axis"][0]}[{role["axis"][1]}] + [{role["bar"]}]','Sort by value'],0,0,0,0),782,326,474,230,9)
   add_visual(pg,textbox(['MATRIX',f'Rows: {role["axis"][1]} / Actual / PY / YoY %','Apply conditional formatting and explain one finding.'],0,0,0,0),24,578,1232,116,10)
  else:
   cards=role['cards'] if idx!=1 else (['Total Revenue','Demo Revenue PM','Demo Revenue MoM %','Demo Revenue PY'] if stage=='START' else ['Total Revenue','Revenue PM','Revenue MoM %','Revenue PY'])
   for j,m in enumerate(cards):add_visual(pg,chart('card',{'Values':[(tab,m,'Measure')]},m),24+j*310,192,296,112,j+4)
   line=chart('lineChart',{'Category':[('DimDate','Year Month','Column')],'Y':[(tab,base,'Measure'),(tab,'Planned Headcount' if tab=='FactWorkforce' else ('Demo Revenue PY' if stage=='START' and idx==1 else prefix+' PY'),'Measure')]},'Actual vs Planned' if tab=='FactWorkforce' else 'Actual vs Previous Year')
   line['query']['sortDefinition']={'sort':[{'field':expression('DimDate','Year Month'),'direction':'Ascending'}]}
   add_visual(pg,line,24,326,740,230,8)
   if stage=='FINAL_REFERENCE' or idx==1:
    bar=chart('barChart',{'Category':[(role['axis'][0],role['axis'][1],'Column')],'Y':[(tab,role['bar'],'Measure')]},role['bar']+' by '+role['axis'][1])
    bar['query']['sortDefinition']={'sort':[{'field':expression(tab,role['bar'],'Measure'),'direction':'Ascending' if tab in ['FactFinance','FactWorkforce'] else 'Descending'}]}
    add_visual(pg,bar,782,326,474,230,9)
    matrix_measures = {'FactSales':([base,'Demo Revenue PY','Demo Revenue MoM %'] if stage=='START' and idx==1 else [base,prefix+' PY',prefix+' YoY %']), 'FactOperations':['Orders','Completed Orders','On Time %','Average Cycle Days'], 'FactFinance':['Finance Revenue','Operating Profit','Budget Variance'], 'FactWorkforce':['Headcount','Planned Headcount','Staffing Gap','Staffing Coverage %']}[tab]
    add_visual(pg,chart('pivotTable',{'Rows':[(role['axis'][0],role['axis'][1],'Column')],'Values':[(tab,m,'Measure') for m in matrix_measures]},'Comparison matrix'),24,578,1232,116,10)
   else:add_visual(pg,textbox(['NEXT: add your bar chart and matrix.','Then test page filters, visual filters and conditional formatting.'],0,0,0,0),782,326,474,230,9)
 write_json(folder/f'Nova_Day07_{stage}.pbip',{'$schema':SCHEMA+'fabric/pbip/pbipProperties/1.0.0/schema.json','version':'1.0','artifacts':[{'report':{'path':'Nova.Report'}}],'settings':{'enableAutoRecovery':True}})

def build_model(folder,stage):
 tables=[]
 for n,(headers,rows) in DATA.items():
  s=io.StringIO();w=csv.writer(s);w.writerow(headers);w.writerows(rows)
  compressor=zlib.compressobj(9,zlib.DEFLATED,-15);encoded=base64.b64encode(compressor.compress(s.getvalue().encode())+compressor.flush()).decode()
  mtypes={'dateTime':'type date','string':'type text','double':'type number','int64':'Int64.Type'}
  typed=', '.join('{"'+h+'", '+mtypes[t]+'}' for h,t in zip(headers,TYPES[n]))
  expr=['let',f'    Source = Csv.Document(Binary.Decompress(Binary.FromText("{encoded}", BinaryEncoding.Base64), Compression.Deflate), [Delimiter=",", Columns={len(headers)}, Encoding=65001, QuoteStyle=QuoteStyle.Csv]),','    Headers = Table.PromoteHeaders(Source, [PromoteAllScalars=true]),',f'    Typed = Table.TransformColumnTypes(Headers, {{{typed}}}, "en-US")','in','    Typed']
  cols=[]
  for h,t in zip(headers,TYPES[n]):
   c={'name':h,'dataType':t,'sourceColumn':h,'summarizeBy':'none' if t in ['string','dateTime'] or n.startswith('Dim') else 'sum'}
   if t=='dateTime':c['formatString']='yyyy-MM-dd';c['annotations']=[{'name':'UnderlyingDateTimeDataType','value':'Date'}]
   if n=='DimDate' and h=='Date':c.update(isKey=True)
   if n=='DimDate' and h=='Month':c['sortByColumn']='Month Number'
   if n=='DimDate' and h=='Year Month':c['sortByColumn']='Year Month Sort'
   cols.append(c)
  measures=[{k:v for k,v in m.items() if k not in ['level','meaning']} for m in MEAS.get(n,[]) if stage!='START' or m['level']=='base']
  if stage=='START' and n=='FactSales':
   for original,demo in [('Revenue PM','Demo Revenue PM'),('Revenue MoM %','Demo Revenue MoM %'),('Revenue PY','Demo Revenue PY')]:
    src=next(m for m in MEAS[n] if m['name']==original)
    measures.append({'name':demo,'expression':src['expression'].replace('[Revenue PM]','[Demo Revenue PM]'),'formatString':src['formatString'],'displayFolder':'Trainer Demo'})
  tab={'name':n,'columns':cols,'partitions':[{'name':n,'mode':'import','source':{'type':'m','expression':expr}}]}
  if measures:tab['measures']=measures
  if n=='DimDate':tab['dataCategory']='Time'
  tables.append(tab)
 rel=[]
 for fact,datecol in [('FactSales','Order Date'),('FactOperations','Order Date'),('FactFinance','Posting Date'),('FactWorkforce','Snapshot Date')]:
  rel.append({'name':f'{fact}_Date','fromTable':fact,'fromColumn':datecol,'toTable':'DimDate','toColumn':'Date','crossFilteringBehavior':'oneDirection','joinOnDateBehavior':'datePartOnly'})
  rel.append({'name':f'{fact}_Branch','fromTable':fact,'fromColumn':'Branch ID','toTable':'DimBranch','toColumn':'Branch ID','crossFilteringBehavior':'oneDirection'})
 for fact,col,dim,dcol in [('FactSales','Product ID','DimProduct','Product ID'),('FactSales','Customer ID','DimCustomer','Customer Code'),('FactWorkforce','Department ID','DimDepartment','Department ID')]:rel.append({'name':f'{fact}_{dim}','fromTable':fact,'fromColumn':col,'toTable':dim,'toColumn':dcol,'crossFilteringBehavior':'oneDirection'})
 model={'name':'Nova_Day07','compatibilityLevel':1567,'model':{'culture':'en-US','defaultPowerBIDataSourceVersion':'powerBI_V3','sourceQueryCulture':'en-US','tables':tables,'relationships':rel,'annotations':[{'name':'PBI_QueryOrder','value':json.dumps(list(DATA))},{'name':'__PBI_TimeIntelligenceEnabled','value':'0'}]}}
 write_json(folder/'Nova.SemanticModel/model.bim',model)
 write_json(folder/'Nova.SemanticModel/definition.pbism',{'$schema':SCHEMA+'fabric/item/semanticModel/definitionProperties/1.0.0/schema.json','version':'1.0','settings':{}})

CSS='''@import url("https://fonts.googleapis.com/css2?family=Almarai:wght@300;400;700;800&family=Inter:wght@400;500;600;700;800&display=swap");*{box-sizing:border-box}body{font:16px/1.7 Almarai,"Segoe UI",sans-serif;color:#13233b;background:#f3f7f7;margin:0}main{max-width:1100px;margin:auto;padding:32px}header{background:#13233b;color:white;padding:28px;border-radius:18px;margin-bottom:22px}h1{margin:0;font-size:30px}h2{color:#087f7f;margin-top:28px}.card{background:white;border-radius:14px;padding:24px;margin:18px 0;border:1px solid #e1e9ea}code,pre{direction:ltr;text-align:left;font-family:Consolas,monospace}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#eef5f5;padding:18px;border-radius:10px}table{width:100%;border-collapse:collapse}td,th{text-align:start;padding:12px;border-bottom:1px solid #e4ebeb}a{color:#087f7f}p{margin:8px 0}small{color:#61717c}.en{direction:ltr;text-align:left}@media(max-width:650px){main{padding:16px}h1{font-size:24px}td,th{padding:7px}}@media print{body{background:white}.card{break-inside:avoid}main{padding:0}header{color:#13233b;background:white;border-bottom:3px solid #099999}}'''
DAX_COPY_ASSETS='<style>.dax-copy{display:block;margin:0 0 8px auto;padding:8px 16px;border:0;border-radius:8px;background:#099999;color:white;font:700 13px Inter,Almarai,sans-serif;cursor:pointer}.dax-copy:focus-visible{outline:3px solid #13233b;outline-offset:3px}.copy-feedback{display:block;min-height:24px;color:#087f7f;font-size:13px}</style><script>\nasync function copyDaxText(text){\n  if(navigator.clipboard && window.isSecureContext){try{await navigator.clipboard.writeText(text);return;}catch(error){}}\n  const field=document.createElement(`textarea`);field.value=text;field.setAttribute(`readonly`,``);field.style.position=`fixed`;field.style.opacity=`0`;document.body.append(field);field.select();field.setSelectionRange(0,field.value.length);\n  let copied=false;try{copied=document.execCommand(`copy`);}finally{field.remove();}if(!copied)throw new Error(`Clipboard unavailable`);\n}\nfor(const pre of document.querySelectorAll(`.card > pre`)){\n  if(pre.previousElementSibling?.tagName!==`H3`)continue;\n  const button=document.createElement(`button`);button.type=`button`;button.className=`dax-copy`;button.textContent=`COPY`;button.setAttribute(`aria-label`,`نسخ الصيغة ${pre.previousElementSibling.textContent}`);\n  const feedback=document.createElement(`span`);feedback.className=`copy-feedback`;feedback.setAttribute(`role`,`status`);feedback.setAttribute(`aria-live`,`polite`);\n  pre.before(button);pre.after(feedback);let reset;\n  button.addEventListener(`click`,async()=>{clearTimeout(reset);try{await copyDaxText(pre.textContent);feedback.textContent=`تم نسخ الصيغة`;button.textContent=`COPIED`;}catch(error){feedback.textContent=`تعذر النسخ. حدد الصيغة وانسخها يدويا.`;}button.focus();reset=setTimeout(()=>{button.textContent=`COPY`;feedback.textContent=``;},2000);});\n}\n</script>'
def document(title,body):return f'<!doctype html><html lang="ar" dir="rtl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)}</title><style>{CSS}</style><main><header><h1>{title}</h1><p class="en">X Academy · Power BI Specialist · Day 7</p></header>{body}</main>{DAX_COPY_ASSETS}</html>'
def build_guides():
 guide=WORK/'03_GUIDES';guide.mkdir(parents=True)
 (guide/'DAX_FOUNDATIONS.html').write_text((ROOT/'week3/day7/web/DAX_FOUNDATIONS.html').read_text().replace('../../../dashboard.html','../00_START_HERE.html'),encoding='utf-8')
 for r in ROLES:
  tab=r['table'];sections=f'<div class="card"><h2>المهمة</h2><p>{r["question"]}</p><p>اختار سنة 2026. للمقارنة الشهرية اختار شهرا واحدا: مارس أو أبريل.</p><p>ابدأ ببطاقات الأرقام، ثم المقارنة، ثم الرسوم. الصيغ الأساسية موجودة في نسخة البداية؛ أنشئ صيغ المقارنة الزمنية الجديدة في جدول القطاع المناسب. احذف صندوق التعليمات من مكان الرسم ثم ضع الرسم مكانه.</p></div>'
  sections+='<div class="card"><h2>ما يجب بناؤه</h2><table><tr><th>العنصر</th><th>الحقول</th></tr>'
  matrix_fields = {'FactSales':'Revenue / PY Revenue / YoY %', 'FactOperations':'Orders / Completed Orders / On Time % / Average Cycle Days', 'FactFinance':'Finance Revenue / Operating Profit / Budget Variance', 'FactWorkforce':'Headcount / Planned Headcount / Staffing Gap / Staffing Coverage %'}[tab]
  line_comparison = 'Planned Headcount' if tab=='FactWorkforce' else r['prefix']+' PY'
  pairs=[('4 KPI Cards',', '.join(r['cards'])),('Line Chart','DimDate[Year Month] / ['+r['base']+'] / ['+line_comparison+']'),('Bar Chart',r['axis'][0]+'['+r['axis'][1]+'] / ['+r['bar']+']'),('Matrix',r['axis'][1]+' / '+matrix_fields),('Slicers','Year / Year Month / '+r['slicer'][1]),('Filters','Page: Year 2026. Visual: select categories or statuses only for the relevant visual.'),('Formatting','Sort bar by value. Add conditional formatting to the comparison matrix. Align cards and use consistent units.')]
  sections+=''.join('<tr><td class="en">'+html.escape(a)+'</td><td class="en">'+html.escape(b)+'</td></tr>' for a,b in pairs)+'</table></div>'
  if tab=='FactWorkforce':sections+='<div class="card"><h2>قاعدة مهمة</h2><p>عدد الموظفين لقطة شهرية. استخدم آخر لقطة متاحة ضمن الفترة، ولا تجمع موظفي يناير وفبراير ومارس. المقارنة السابقة تعني آخر لقطة في الفترة السابقة.</p><p>لا نستخدم مجموع الموظفين منذ بداية السنة. قياس العجز سالب عند نقص الموظفين.</p></div>'
  if tab=='FactOperations':sections+='<div class="card"><h2>تعريف الالتزام</h2><p>عدد الطلبات المنجزة ضمن المهلة مقسوما على جميع الطلبات المنجزة فقط. الطلبات المفتوحة والملغاة خارج المقام. الاتجاه يعتمد تاريخ إنشاء الطلب.</p><p>الفرق بين نسبتين يعرض كنقاط مئوية؛ ارتفاع النسبة من 80% إلى 85% يعني 5 نقاط مئوية.</p></div>'
  image=r['name'].lower()+'-'+r['domain'].lower()
  reference=f'../05_VISUAL_REFERENCES/{image}.webp'
  sections=f'<div class="card"><h2>اللوحة المتوقعة — {r["ar"]}</h2><p>مرجع للشكل النهائي المطلوب. اضغط على الصورة لعرضها بالحجم الكامل.</p><a href="{reference}" target="_blank" rel="noopener"><img src="{reference}" alt="اللوحة المتوقعة — {r["ar"]}" style="display:block;width:100%;height:auto;border-radius:12px" decoding="async"></a><p><a href="{reference}" download>تنزيل صورة اللوحة</a></p></div>'+sections
  sections+='<div class="card"><h2>قبل كتابة الصيغ</h2><p><a href="DAX_FOUNDATIONS.html">افهم الدوال الجديدة في دليل شرح الصيغ ↗</a></p></div>'
  sections+='<h2>الصيغ المطلوبة</h2>'
  for m in MEAS[tab]:
   status='موجود في نسخة البداية' if m['level']=='base' else 'أنشئه أثناء التطبيق'
   sections+=f'<div class="card"><p>{status}</p><h3 class="en">{html.escape(m["name"])}</h3><pre>{html.escape(m["name"]+" =\n"+m["expression"])}</pre><small class="en">Format: {html.escape(m["formatString"])}'+(' · '+html.escape(m['meaning']) if m['meaning'] else '')+'</small></div>'
  sections+='<div class="card"><h2>نقاط التحقق</h2><p>15 دقيقة: بطاقات الأرقام تستجيب لفلتر السنة.</p><p>30 دقيقة: المقارنة الشهرية تتغير عند اختيار مارس وأبريل، والفترة الأولى تظهر بدون مقارنة سابقة.</p><p>45 دقيقة: الاتجاه والرسم والأرقام متسقة، وفرز الأشهر صحيح.</p><p>النهاية: فلتر صفحة وفلتر رسم، تنسيق شرطي، واستنتاج واحد مرتبط برقم. تحقق من إحدى النتائج مع المدرب أثناء المحاضرة.</p><p>إذا تعطلت: اطلب مساعدة المدرب أثناء المحاضرة، وراجع معه الصيغة أو الرسم ثم واصل التطبيق.</p></div>'
  curated=ROOT/'week3/day7/web'/f'{r["name"]}_{r["domain"]}.html'
  content=curated.read_text().replace('../images/','../05_VISUAL_REFERENCES/').replace('../icons/','../06_ICONS/').replace('.png','.webp')
  (guide/curated.name).write_text(content,encoding='utf-8')
  (guide/f'{r["name"]}_DAX.txt').write_text('\n\n'.join(m['name']+' =\n'+m['expression'] for m in MEAS[tab]),encoding='utf-8')
 sources='''<div class="card"><h2>بداية التشغيل</h2><p>افتح ملف البداية مباشرة داخل باور بي آي، ثم انتقل إلى الصفحة التي تحمل اسمك واتبع دليل التدريب. البيانات مدمجة داخل الملف ولا تحتاج إرفاق ملفات المصدر لفتحه.</p><pre>02_PROJECTS/START/Nova_Day07_START.pbix</pre><p>للتحديث والتحقق قبل المحاضرة:</p><pre>Home → Refresh
Table tools → Mark as date table → DimDate[Date]</pre><p>اختر سنة 2026. استخدم إصدارا حديثا من البرنامج؛ التحويل لا يضمن التوافق مع الإصدارات القديمة.</p><p>ملف البداية محول ومحفوظ بواسطة المدرب. فحصنا سلامة الحزمة والصفحات الست ووجود نموذج البيانات؛ اختبار الحسابات والتحديث داخل البرنامج يبقى ضمن مراجعة المدرب.</p><p>نسختا المتابعة والمرجع متاحتان للمدرب في حزمة منفصلة بصيغة المشروع، وتحتاجان الفتح والتحويل داخل البرنامج.</p><pre>Day07_TRAINER_REFERENCES.zip</pre></div>'''
 sources+='''<div class="card"><h2>البيانات</h2><p>البيانات تدريبية. حافظنا على ملف اليوم الرابع كاملا دون تغيير، وعلى جميع معاملات المبيعات الأصلية البالغ عددها 40,000. أضفنا 40,000 معاملة تدريبية للمقارنة بسنة 2025.</p><p>كل القطاعات تغطي يناير إلى أبريل في 2025 و2026. جدول التاريخ يغطي السنتين كاملتين. الأشهر بدون بيانات لا تمثل مبيعات صفرا مثبتة.</p><p>العمليات: صف واحد لكل طلب. المالية: صف واحد لكل فرع وشهر، بقيم الإيراد والتكلفة وموازنة الإيراد بالدينار الأردني. الموارد: لقطة واحدة لكل فرع وقسم في نهاية الشهر.</p><p>المالية والعمليات والموارد حالات مستقلة تستخدم أبعادا مشتركة؛ لا يشترط أن تتطابق أرقام إيراد المالية مع معاملات مبيعات اليوم الرابع.</p></div>'''
 sources+='''<div class="card"><h2>خطة المدرب — ساعتان</h2><table><tr><th>الوقت</th><th>التطبيق</th></tr><tr><td>0–10</td><td>فتح المشروع وتحديث البيانات والتأكد من العلاقات.</td></tr><tr><td>10–25</td><td>مثال مشترك: إيراد الشهر السابق ونسبة التغير، ثم الفرق بين فلتر الصفحة وفلتر الرسم.</td></tr><tr><td>25–40</td><td>كل متدرب يبني بطاقاته. أول نقطة تحقق.</td></tr><tr><td>40–55</td><td>مقارنة الشهر السابق أو السنة الماضية. ثاني نقطة تحقق.</td></tr><tr><td>55–70</td><td>الاتجاه والرسم والفرز. ثالث نقطة تحقق.</td></tr><tr><td>70–95</td><td>الفلاتر والمصفوفة والتنسيق الشرطي.</td></tr><tr><td>95–110</td><td>اختبار الأرقام ومقارنة النتائج بالمرجع.</td></tr><tr><td>110–120</td><td>استنتاج واحد لكل متدرب وحفظ الملف.</td></tr></table></div>'''
 sources+='''<div class="card"><h2>المتابعة والمرجع</h2><p>البداية: النموذج والصيغ الأساسية، صفحة المثال المشترك، وأماكن واضحة لكل رسم.</p><p>المتابعة: الصيغ الزمنية وبطاقات الأرقام والاتجاه لكل متدرب؛ يكمل الرسوم والمصفوفة بنفسه.</p><p>المرجع: الصيغ وبطاقات الأرقام والاتجاه والرسم والمصفوفة. يضيف المدرب فلتر السنة والتنسيق الشرطي بعد التحقق داخل البرنامج.</p><p>ابدأ بالسنة 2026 كاملة لرؤية اتجاه يناير إلى أبريل. اختر شهرا واحدا لقياس التغير الشهري؛ اختيار أكثر من شهر يعيد قياس الشهر السابق فارغا عمدا.</p></div>'''
 sources+='<div class="card"><h2>التصاميم المرجعية</h2><p><a href="05_VISUAL_REFERENCES/gallery.html">اللوحات الأربع المتوقعة</a></p></div>'
 sources+='<div class="card"><h2>أدلة المتدربين</h2>'+''.join(f'<p><a href="03_GUIDES/{r["name"]}_{r["domain"]}.html">{r["ar"]} — {r["domain"]}</a></p>' for r in ROLES)+'</div>'
 sources+='''<div class="card"><h2>الانتقال إلى الأيام القادمة</h2><p>اليوم الثامن: نكمل نفس اللوحات بتجربة الاستخدام، التفاعل، صفحات التفاصيل، التلميحات، الأزرار والإشارات المرجعية.</p><p>اليوم التاسع: نربط النتيجة بمحركاتها، ونقدم ثلاثة استنتاجات وإجراء واحد مناسب لدور كل متدرب.</p><p>اليوم العاشر: تحديث البيانات والتحقق ثم عرض نهائي يدافع فيه كل متدرب عن قراره.</p></div>'''
 sources+='''<div class="card"><h2>المرجع التقني</h2><p class="en"><a href="https://learn.microsoft.com/en-us/power-bi/developer/projects/projects-overview">Microsoft Learn — Power BI Desktop projects</a></p><p class="en"><a href="https://learn.microsoft.com/en-us/power-bi/developer/projects/projects-report">Microsoft Learn — Report definitions</a></p><p class="en"><a href="https://learn.microsoft.com/en-us/power-bi/developer/projects/projects-dataset">Microsoft Learn — Semantic model definitions</a></p></div>'''
 (WORK/'00_START_HERE.html').write_text(document('اليوم السابع — أربع لوحات وأربع قرارات',sources),encoding='utf-8')
 # Web copies use relative guide links within their own published folder.
 web=ROOT/'week3/day7/web';web.mkdir(parents=True,exist_ok=True)
 (web/'index.html').write_text(document('اليوم السابع — أربع لوحات وأربع قرارات',sources.replace('03_GUIDES/','').replace('05_VISUAL_REFERENCES/gallery.html','gallery.html')),encoding='utf-8')
 for p in guide.glob('*.html'):(web/p.name).write_text(p.read_text().replace('../05_VISUAL_REFERENCES/','../images/').replace('../06_ICONS/','../icons/').replace('.webp','.png'),encoding='utf-8')


def validate(original):
 column_names={c for headers,rows in DATA.values() for c in headers}
 measure_names=[m['name'] for ms in MEAS.values() for m in ms]
 assert len(measure_names)==len(set(measure_names)), 'Duplicate measure name'
 assert not column_names.intersection(measure_names), 'Column/measure name conflict: '+str(column_names.intersection(measure_names))
 # Static checks and independent numeric controls. This does not assert Desktop execution.
 assert (WORK/'00_DAY04_ORIGINAL/Nova_Day04_Model_Lab.xlsx').read_bytes()==original
 for n,(h,rows) in DATA.items():
  assert all(len(r)==len(h) for r in rows),n
 dims={n:{r[0] for r in rows} for n,(h,rows) in DATA.items() if n.startswith('Dim')}
 for n,keys in dims.items():assert len(keys)==len(DATA[n][1]),n
 for n in ['FactSales','FactOperations','FactFinance','FactWorkforce']:
  h,rows=DATA[n]
  for c,dim in [('Branch ID','DimBranch'),('Product ID','DimProduct'),('Customer ID','DimCustomer'),('Department ID','DimDepartment')]:
   if c in h:assert all(r[h.index(c)] in dims[dim] for r in rows),(n,c)
  datecol={'FactSales':'Order Date','FactOperations':'Order Date','FactFinance':'Posting Date','FactWorkforce':'Snapshot Date'}[n]
  assert all(r[h.index(datecol)] in dims['DimDate'] for r in rows)
 for n,keycols in [('FactSales',['Transaction ID']),('FactOperations',['Order ID']),('FactFinance',['Record ID']),('FactWorkforce',['Snapshot Date','Branch ID','Department ID'])]:
  h,rs=DATA[n];assert len({tuple(r[h.index(c)] for c in keycols) for r in rs})==len(rs),n
 rows=DATA['FactSales'][1]
 assert len(rows)==80000
 assert sum(r[-1]=='Day04 original' for r in rows)==40000
 controls={'data_rows':{n:len(r) for n,(h,r) in DATA.items()},'original_day4_sha256':hashlib.sha256(original).hexdigest(),'desktop_tested':False,'period':'Jan-Apr 2025 and Jan-Apr 2026','latest_month_checks':{}}
 for year in [2025,2026]:
  for mo in range(1,5):
   ym=f'{year}-{mo:02}';s=[r for r in DATA['FactSales'][1] if r[1].startswith(ym)]
   o=[r for r in DATA['FactOperations'][1] if r[1].startswith(ym)];f=[r for r in DATA['FactFinance'][1] if r[1].startswith(ym)];hr=[r for r in DATA['FactWorkforce'][1] if r[0].startswith(ym)]
   done=[r for r in o if r[3]=='Completed']
   controls['latest_month_checks'][ym]={'sales_revenue':round(sum(r[8] for r in s),2),'sales_transactions':len(s),'sales_customers':len({r[2] for r in s}),'orders':len(o),'completed_orders':len(done),'on_time_pct':sum(r[6] for r in done)/len(done),'finance_revenue':round(sum(r[3] for r in f),2),'operating_profit':round(sum(r[3]-r[4] for r in f),2),'headcount':sum(r[4] for r in hr),'planned_headcount':sum(r[3] for r in hr)}
 write_json(WORK/'04_CHECKS/Expected_Results.json',controls)
 h=['Period']+list(next(iter(controls['latest_month_checks'].values())))
 p=WORK/'04_CHECKS/Expected_Results.csv'
 with p.open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.writer(f);w.writerow(h)
  for ym,vals in controls['latest_month_checks'].items():w.writerow([ym]+list(vals.values()))
 (WORK/'04_CHECKS/Desktop_Checklist.txt').write_text('BEFORE CLASS\n1. Open Nova_Day07_START.pbix in current Power BI Desktop.\n2. Refresh and confirm all 9 tables load.\n3. Confirm 11 active single-direction 1:* relationships.\n4. Mark DimDate as Date table using Date.\n5. Select Year 2026 and Year Month 2026-04.\n6. Compare cards to Expected_Results.csv.\n7. Check March and April PM and MoM results.\n8. HR total headcount must use the last snapshot, not the sum of months.\n9. Clear month slicer and select Year 2026; inspect Jan-Apr trend.\n10. Download Day07_TRAINER_REFERENCES.zip and open CHECKPOINT and FINAL_REFERENCE; confirm every visual binds correctly.\n11. Set default page Year filter 2026 and add requested conditional formatting.\n12. Save each as PBIX and distribute the tested copies.\n\nSTATIC VALIDATION DOES NOT REPLACE THESE DESKTOP CHECKS.\n',encoding='utf-8')
 return controls

def main():
 if WORK.exists():shutil.rmtree(WORK)
 WORK.mkdir(parents=True)
 original=build_data();build_measures()
 for stage in ['START','CHECKPOINT','FINAL_REFERENCE']:
  folder=WORK/'02_PROJECTS'/stage;build_model(folder,stage);build_report(folder,stage)
 build_guides();shutil.copytree(ROOT/'week3/day7/icons',WORK/'06_ICONS',dirs_exist_ok=True);controls=validate(original)
 visuals=WORK/'05_VISUAL_REFERENCES';visuals.mkdir(parents=True,exist_ok=True)
 for image in (ROOT/'week3/day7/images').glob('*.png'):
  Image.open(image).save(visuals/(image.stem+'.webp'),format='WEBP',quality=96,method=6)
 gallery=(ROOT/'week3/day7/web/gallery.html').read_text().replace('../images/','').replace('.png','.webp').replace('href="index.html"','href="../00_START_HERE.html"')
 for role in ROLES:gallery=gallery.replace('href="'+role['name']+'_'+role['domain']+'.html"','href="../03_GUIDES/'+role['name']+'_'+role['domain']+'.html"')
 (visuals/'gallery.html').write_text(gallery,encoding='utf-8')
 theme={'name':'X Academy Nova','dataColors':['#099999','#426A8F','#B99355','#7966A8','#658977'],'background':'#F3F7F7','foreground':'#13233B','tableAccent':'#099999','textClasses':{'title':{'fontFace':'Segoe UI','fontSize':14},'label':{'fontFace':'Segoe UI','fontSize':11}}}
 write_json(WORK/'03_GUIDES/X_Academy_Nova_Theme.json',theme)
 pbix=ROOT/'week3/day7/downloads/Nova_Day07_START.pbix'
 if not pbix.exists():raise FileNotFoundError('The trainer-converted START PBIX is required.')
 with zipfile.ZipFile(pbix) as check:
  if check.testzip() is not None or 'DataModel' not in check.namelist():raise ValueError('Invalid START PBIX')
 with zipfile.ZipFile(OUT.with_name('Day07_TRAINER_REFERENCES.zip'),'w',zipfile.ZIP_DEFLATED,compresslevel=9) as refs:
  for stage in ['CHECKPOINT','FINAL_REFERENCE']:
   for p in sorted((WORK/'02_PROJECTS'/stage).rglob('*')):
    if p.is_file():refs.write(p,p.relative_to(WORK))
 shutil.rmtree(WORK/'02_PROJECTS')
 starter=WORK/'02_PROJECTS/START';starter.mkdir(parents=True)
 shutil.copy2(pbix,starter/pbix.name)
 OUT.parent.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
  for p in sorted(WORK.rglob('*')):
   if p.is_file():z.write(p,p.relative_to(WORK))
 print(json.dumps({'zip':str(OUT),'bytes':OUT.stat().st_size,'rows':controls['data_rows'],'april2026':controls['latest_month_checks']['2026-04']},indent=2))
if __name__=='__main__':main()
